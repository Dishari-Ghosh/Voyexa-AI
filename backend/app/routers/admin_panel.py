"""
Everything the admin website (/admin on the frontend) talks to.

  PUBLIC   POST /admin/login    - admin-only login (normal users are refused)
           POST /admin/signup   - create an admin account; needs ADMIN_SIGNUP_KEY

  ADMIN    GET    /admin/stats                dashboard numbers
  ONLY     GET    /admin/users                user management (read-only)
           GET    /admin/feedback             feedback management
           PATCH  /admin/feedback/{id}        mark reviewed / not reviewed
           DELETE /admin/feedback/{id}
           GET    /admin/places               place management (list/search)
           POST   /admin/places               add a place
           PUT    /admin/places/{id}          update a place
           DELETE /admin/places/{id}          remove a place
           POST   /admin/places/{id}/restore  bring a removed place back
           GET    /admin/analytics            chart data
           GET    /admin/export/{dataset}     Excel download
           GET/PUT /admin/settings            Power BI embed link

Every protected route requires an account whose role is "admin" - a normal
user's token gets a 403 here, no matter what the frontend shows.
(The older /admin/images routes still live in admin.py.)
"""
import io
import re
import secrets
from collections import Counter
from datetime import date, datetime, timezone
from typing import Optional

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr, Field

from app.auth.dependencies import require_admin
from app.auth.security import create_access_token, hash_password, verify_password
from app.config import settings
from app.db.mongo import (
    feedback_collection,
    place_overrides_collection,
    settings_collection,
    trips_collection,
    users_collection,
    wishlist_collection,
)
from app.models.user import Token, UserLogin
from app.routers.auth import _to_user_out
from app.services import place_overlay as po

# Routes anyone can reach (they only ever hand out a token to a real admin).
public_router = APIRouter(prefix="/admin", tags=["admin-panel"])
# Routes that need an admin token.
router = APIRouter(prefix="/admin", tags=["admin-panel"], dependencies=[Depends(require_admin)])

DEFAULT_ADMIN_KEYS = {"", "change-this-too"}


# --------------------------------------------------------------------------
# Admin login / signup
# --------------------------------------------------------------------------
class AdminSignup(BaseModel):
    name: str
    email: EmailStr
    password: str = Field(min_length=8)
    phone: Optional[str] = None
    dob: Optional[date] = None
    admin_key: str


@public_router.post("/login", response_model=Token)
async def admin_login(payload: UserLogin):
    doc = await users_collection().find_one({"email": payload.email})
    if not doc or not verify_password(payload.password, doc["password_hash"]):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Incorrect email or password")
    if doc.get("role") != "admin":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This account doesn't have admin access.")
    token = create_access_token({"sub": str(doc["_id"])})
    return Token(access_token=token, user=_to_user_out(doc))


@public_router.post("/signup", response_model=Token)
async def admin_signup(payload: AdminSignup):
    configured = settings.ADMIN_SIGNUP_KEY
    if configured in DEFAULT_ADMIN_KEYS:
        raise HTTPException(
            status.HTTP_503_SERVICE_UNAVAILABLE,
            "Admin signup is switched off until ADMIN_SIGNUP_KEY is set to a private value on the server.",
        )
    if not secrets.compare_digest(payload.admin_key.encode(), configured.encode()):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "That admin key is not correct.")

    existing = await users_collection().find_one({"email": payload.email})
    if existing:
        # Someone who already has a normal account can be promoted, but only
        # by proving they own it (correct password) AND knowing the key.
        if not verify_password(payload.password, existing["password_hash"]):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Email already registered")
        await users_collection().update_one({"_id": existing["_id"]}, {"$set": {"role": "admin"}})
        existing["role"] = "admin"
        doc = existing
    else:
        doc = {
            "name": payload.name,
            "email": payload.email,
            "password_hash": hash_password(payload.password),
            "phone": payload.phone,
            "dob": payload.dob.isoformat() if payload.dob else None,
            "role": "admin",
        }
        doc["_id"] = (await users_collection().insert_one(doc)).inserted_id

    token = create_access_token({"sub": str(doc["_id"])})
    return Token(access_token=token, user=_to_user_out(doc))


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------
@router.get("/stats")
async def stats():
    ov = await po.load_overlay()
    places = po.effective_places(ov)

    total_feedback = await feedback_collection().count_documents({})
    new_feedback = await feedback_collection().count_documents({"reviewed": {"$ne": True}})

    recent_feedback = [
        _feedback_out(d) async for d in feedback_collection().find({}).sort("created_at", -1).limit(5)
    ]
    recent_users = [
        _user_row(d) async for d in users_collection().find({}).sort("_id", -1).limit(5)
    ]

    return {
        "total_users": await users_collection().count_documents({"role": {"$ne": "admin"}}),
        "total_admins": await users_collection().count_documents({"role": "admin"}),
        "total_places": len(places),
        "total_trips": await trips_collection().count_documents({}),
        "total_feedback": total_feedback,
        "new_feedback": new_feedback,
        "hidden_gems": sum(1 for p in places if p.get("Is_Hidden_Gem")),
        "recent_feedback": recent_feedback,
        "recent_users": recent_users,
    }


# --------------------------------------------------------------------------
# User management (read-only)
# --------------------------------------------------------------------------
def _user_row(doc: dict) -> dict:
    return {
        "id": str(doc["_id"]),
        "name": doc.get("name", ""),
        "email": doc.get("email", ""),
        "phone": doc.get("phone"),
        "dob": doc.get("dob"),
        "role": doc.get("role", "user"),
        # every MongoDB id carries its creation time, so no extra field is needed
        "joined": doc["_id"].generation_time.isoformat(),
    }


@router.get("/users")
async def list_users(
    search: str = "",
    role: str = "",
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
):
    query: dict = {}
    if role in ("user", "admin"):
        query["role"] = role if role == "admin" else {"$ne": "admin"}
    if search.strip():
        rx = {"$regex": re.escape(search.strip()), "$options": "i"}
        query["$or"] = [{"name": rx}, {"email": rx}, {"phone": rx}]

    total = await users_collection().count_documents(query)
    cursor = (
        users_collection().find(query).sort("_id", -1).skip((page - 1) * page_size).limit(page_size)
    )
    return {"total": total, "page": page, "page_size": page_size, "users": [_user_row(d) async for d in cursor]}


# --------------------------------------------------------------------------
# Feedback management
# --------------------------------------------------------------------------
def _feedback_out(doc: dict) -> dict:
    return {
        "id": str(doc["_id"]),
        "name": doc.get("name") or "Anonymous",
        "rating": doc.get("rating", 0),
        "message": doc.get("message", ""),
        "created_at": doc.get("created_at"),
        "reviewed": bool(doc.get("reviewed", False)),
        "reviewed_at": doc.get("reviewed_at"),
    }


class ReviewUpdate(BaseModel):
    reviewed: bool


def _oid(value: str) -> ObjectId:
    try:
        return ObjectId(value)
    except Exception:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Not found")


@router.get("/feedback")
async def list_feedback(status_filter: str = Query("all", alias="status")):
    query: dict = {}
    if status_filter == "new":
        query["reviewed"] = {"$ne": True}
    elif status_filter == "reviewed":
        query["reviewed"] = True
    items = [_feedback_out(d) async for d in feedback_collection().find(query).sort("created_at", -1)]
    return {
        "feedback": items,
        "total": await feedback_collection().count_documents({}),
        "new": await feedback_collection().count_documents({"reviewed": {"$ne": True}}),
    }


@router.patch("/feedback/{feedback_id}")
async def mark_feedback(feedback_id: str, body: ReviewUpdate):
    update = {"reviewed": body.reviewed, "reviewed_at": datetime.now(timezone.utc).isoformat() if body.reviewed else None}
    res = await feedback_collection().update_one({"_id": _oid(feedback_id)}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Feedback not found")
    doc = await feedback_collection().find_one({"_id": _oid(feedback_id)})
    return _feedback_out(doc)


@router.delete("/feedback/{feedback_id}")
async def delete_feedback(feedback_id: str):
    await feedback_collection().delete_one({"_id": _oid(feedback_id)})
    return {"deleted": True}


# --------------------------------------------------------------------------
# Place management
# --------------------------------------------------------------------------
class PlaceFields(BaseModel):
    """Used for both add (name/city/state/zone/lat/lng required by the route)
    and update (any subset)."""
    Place_Name: Optional[str] = Field(None, min_length=1, max_length=120)
    City: Optional[str] = Field(None, max_length=80)
    State: Optional[str] = Field(None, max_length=80)
    Zone: Optional[str] = Field(None, max_length=20)
    Place_Type: Optional[str] = Field(None, max_length=60)
    Interest_Tags: Optional[str] = Field(None, max_length=200)
    Suitable_For: Optional[str] = Field(None, max_length=200)
    Entry_Fee_INR: Optional[int] = Field(None, ge=0)
    Activity_Cost_Min: Optional[int] = Field(None, ge=0)
    Activity_Cost_Max: Optional[int] = Field(None, ge=0)
    HotelcostpernightINR_Min: Optional[int] = Field(None, ge=0)
    HotelcostpernightINR_Max: Optional[int] = Field(None, ge=0)
    FoodcostperdayINR_Min: Optional[int] = Field(None, ge=0)
    FoodcostperdayINR_Max: Optional[int] = Field(None, ge=0)
    Latitude: Optional[float] = Field(None, ge=-90, le=90)
    Longitude: Optional[float] = Field(None, ge=-180, le=180)
    Popularity_Rating: Optional[float] = Field(None, ge=0, le=5)
    Best_Month_Start: Optional[int] = Field(None, ge=1, le=12)
    Best_Month_End: Optional[int] = Field(None, ge=1, le=12)
    Terrain_Type: Optional[str] = Field(None, max_length=60)
    Difficulty_Level: Optional[str] = Field(None, max_length=30)
    Typical_Duration: Optional[str] = Field(None, max_length=60)
    Is_Hidden_Gem: Optional[bool] = None
    Hidden_Gem_Score: Optional[float] = Field(None, ge=0, le=100)


def _decorate(p: dict, ov: po.Overlay, custom_ids: set[int]) -> dict:
    pid = p["place_id"]
    return {
        **p,
        "_custom": pid in custom_ids,
        "_edited": pid in ov.edits,
        "_removed": pid in ov.deleted,
    }


@router.get("/places")
async def list_places(
    search: str = "",
    state: str = "",
    view: str = Query("live", pattern="^(live|removed)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    ov = await po.load_overlay()
    custom_ids = {c["place_id"] for c in ov.custom}

    if view == "removed":
        merged = [{**p, **ov.edits.get(p["place_id"], {})} for p in po.base_places() if p["place_id"] in ov.deleted]
    else:
        merged = po.effective_places(ov)

    if state:
        merged = [p for p in merged if p.get("State") == state]
    if search.strip():
        q = search.strip().lower()
        merged = [
            p for p in merged
            if q in str(p.get("Place_Name", "")).lower()
            or q in str(p.get("City", "")).lower()
            or q in str(p.get("State", "")).lower()
            or q == str(p.get("place_id"))
        ]

    # newest admin-added places first, then by id
    merged.sort(key=lambda p: (0 if p["place_id"] in custom_ids else 1, -p["place_id"] if p["place_id"] in custom_ids else p["place_id"]))

    total = len(merged)
    start = (page - 1) * page_size
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "places": [_decorate(p, ov, custom_ids) for p in merged[start:start + page_size]],
        "removed_count": len(ov.deleted),
    }


@router.get("/places/options")
async def place_options():
    """Existing states / zones / types, to fill the dropdowns in the place form."""
    ov = await po.load_overlay()
    places = po.effective_places(ov)
    def uniq(key): return sorted({str(p.get(key)) for p in places if p.get(key)})
    return {
        "states": uniq("State"),
        "zones": uniq("Zone"),
        "place_types": uniq("Place_Type"),
        "terrains": uniq("Terrain_Type"),
        "difficulties": uniq("Difficulty_Level"),
        "durations": uniq("Typical_Duration"),
    }


@router.post("/places")
async def add_place(body: PlaceFields):
    data = po.coerce_fields(body.model_dump(exclude_none=True))
    missing = [k for k in ("Place_Name", "City", "State", "Zone", "Latitude", "Longitude") if k not in data or data[k] == ""]
    if missing:
        raise HTTPException(422, f"Please fill in: {', '.join(missing)}")

    full = {**po.CUSTOM_DEFAULTS, **data}
    # keep each min/max pair sensible
    for lo, hi in (("Activity_Cost_Min", "Activity_Cost_Max"),
                   ("HotelcostpernightINR_Min", "HotelcostpernightINR_Max"),
                   ("FoodcostperdayINR_Min", "FoodcostperdayINR_Max")):
        if full[hi] < full[lo]:
            full[hi] = full[lo]

    pid = await po.next_custom_id()
    await place_overrides_collection().insert_one(
        {"place_id": pid, "custom": True, "deleted": False, "fields": full}
    )
    return {"place_id": pid, **full, "_custom": True}


@router.put("/places/{place_id}")
async def update_place(place_id: int, body: PlaceFields):
    changes = po.coerce_fields(body.model_dump(exclude_none=True))
    if not changes:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Nothing to update")

    coll = place_overrides_collection()
    existing = await coll.find_one({"place_id": place_id})

    if existing and existing.get("custom"):
        fields = {**existing.get("fields", {}), **changes}
        await coll.update_one({"place_id": place_id}, {"$set": {"fields": fields}})
        return {"place_id": place_id, **fields, "_custom": True}

    if not any(p["place_id"] == place_id for p in po.base_places()):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Place not found")

    fields = {**(existing or {}).get("fields", {}), **changes}
    await coll.update_one(
        {"place_id": place_id},
        {"$set": {"place_id": place_id, "fields": fields}, "$setOnInsert": {"deleted": False}},
        upsert=True,
    )
    base = next(p for p in po.base_places() if p["place_id"] == place_id)
    return {**base, **fields, "_edited": True}


@router.delete("/places/{place_id}")
async def remove_place(place_id: int):
    coll = place_overrides_collection()
    existing = await coll.find_one({"place_id": place_id})
    if existing and existing.get("custom"):
        await coll.delete_one({"place_id": place_id})  # admin-added: gone for good
        return {"removed": True, "permanent": True}

    if not any(p["place_id"] == place_id for p in po.base_places()):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Place not found")
    await coll.update_one(
        {"place_id": place_id},
        {"$set": {"place_id": place_id, "deleted": True}},
        upsert=True,
    )
    return {"removed": True, "permanent": False}


@router.post("/places/{place_id}/restore")
async def restore_place(place_id: int):
    await place_overrides_collection().update_one({"place_id": place_id, "custom": {"$ne": True}}, {"$set": {"deleted": False}})
    return {"restored": True}


# --------------------------------------------------------------------------
# Analytics
# --------------------------------------------------------------------------
def _month_key(dt: datetime) -> str:
    return f"{dt.year}-{dt.month:02d}"


def _last_12_months() -> list[str]:
    now = datetime.now(timezone.utc)
    y, m = now.year, now.month
    keys = []
    for _ in range(12):
        keys.append(f"{y}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return list(reversed(keys))


def _counter_rows(counter: Counter, limit: Optional[int] = None) -> list[dict]:
    items = counter.most_common(limit)
    return [{"label": str(k), "value": v} for k, v in items]


@router.get("/analytics")
async def analytics():
    ov = await po.load_overlay()
    places = po.effective_places(ov)
    names = {p["place_id"]: p.get("Place_Name", str(p["place_id"])) for p in places}

    months = _last_12_months()
    signups = Counter()
    async for d in users_collection().find({}, {"_id": 1}):
        signups[_month_key(d["_id"].generation_time)] += 1

    trips_by_month = Counter()
    planned = Counter()
    async for t in trips_collection().find({}, {"saved_at": 1, "place_ids": 1}):
        try:
            trips_by_month[_month_key(datetime.fromisoformat(t["saved_at"].replace("Z", "+00:00")))] += 1
        except Exception:
            pass
        for pid in t.get("place_ids", []):
            planned[pid] += 1

    saved = Counter()
    async for w in wishlist_collection().find({}, {"place_id": 1}):
        saved[w["place_id"]] += 1

    ratings = Counter()
    async for f in feedback_collection().find({}, {"rating": 1}):
        ratings[int(f.get("rating", 0) or 0)] += 1

    by_zone = Counter(p.get("Zone") for p in places if p.get("Zone"))
    by_state = Counter(p.get("State") for p in places if p.get("State"))
    by_type = Counter(p.get("Place_Type") for p in places if p.get("Place_Type"))
    gems_by_zone = Counter(p.get("Zone") for p in places if p.get("Is_Hidden_Gem") and p.get("Zone"))

    return {
        "signups_by_month": [{"label": m, "value": signups.get(m, 0)} for m in months],
        "trips_by_month": [{"label": m, "value": trips_by_month.get(m, 0)} for m in months],
        "feedback_ratings": [{"label": f"{r} star" + ("" if r == 1 else "s"), "value": ratings.get(r, 0)} for r in range(5, 0, -1)],
        "places_by_zone": _counter_rows(by_zone),
        "places_by_state": _counter_rows(by_state, 10),
        "places_by_type": _counter_rows(by_type, 10),
        "hidden_gems_by_zone": _counter_rows(gems_by_zone),
        "most_planned_places": [{"label": names.get(pid, f"Place {pid}"), "value": n} for pid, n in planned.most_common(10)],
        "most_saved_places": [{"label": names.get(pid, f"Place {pid}"), "value": n} for pid, n in saved.most_common(10)],
    }


# --------------------------------------------------------------------------
# Excel export  (openpyxl is already in requirements.txt)
# --------------------------------------------------------------------------
PLACE_COLUMNS = [
    "place_id", "Place_Name", "City", "State", "Zone", "Place_Type", "Interest_Tags", "Suitable_For",
    "Entry_Fee_INR", "Activity_Cost_Min", "Activity_Cost_Max", "HotelcostpernightINR_Min",
    "HotelcostpernightINR_Max", "FoodcostperdayINR_Min", "FoodcostperdayINR_Max", "Latitude", "Longitude",
    "Popularity_Rating", "Best_Month_Start", "Best_Month_End", "Terrain_Type", "Difficulty_Level",
    "Typical_Duration", "Hidden_Gem_Score", "Is_Hidden_Gem",
]


async def _sheet_rows(dataset: str) -> tuple[list[str], list[list]]:
    if dataset == "users":
        header = ["Name", "Email", "Phone", "Date of birth", "Role", "Joined"]
        rows = []
        async for d in users_collection().find({}).sort("_id", -1):
            u = _user_row(d)
            rows.append([u["name"], u["email"], u["phone"] or "", u["dob"] or "", u["role"], u["joined"][:10]])
        return header, rows

    if dataset == "feedback":
        header = ["Name", "Rating", "Message", "Date given", "Reviewed", "Reviewed at"]
        rows = []
        async for d in feedback_collection().find({}).sort("created_at", -1):
            f = _feedback_out(d)
            rows.append([f["name"], f["rating"], f["message"], (f["created_at"] or "")[:19].replace("T", " "),
                         "Yes" if f["reviewed"] else "No", (f["reviewed_at"] or "")[:19].replace("T", " ")])
        return header, rows

    if dataset == "trips":
        ov = await po.load_overlay()
        by_id = {p["place_id"]: p for p in po.effective_places(ov)}
        header = ["Trip saved at", "Trip for", "Suitable for", "Days", "Budget per day (INR)", "Month", "Interests",
                  "Number of places", "States", "Places"]
        rows = []
        async for t in trips_collection().find({}).sort("saved_at", -1):
            form = t.get("form", {}) or {}
            ps = [by_id[i] for i in t.get("place_ids", []) if i in by_id]
            rows.append([
                (t.get("saved_at") or "")[:19].replace("T", " "),
                form.get("tripFor", ""), form.get("suitableFor", ""), form.get("numDays", ""),
                form.get("budgetPerDay", ""), form.get("month", ""), ", ".join(form.get("interests", []) or []),
                len(t.get("place_ids", [])), ", ".join(sorted({p.get("State", "") for p in ps})),
                ", ".join(p.get("Place_Name", "") for p in ps),
            ])
        return header, rows

    if dataset == "places":
        ov = await po.load_overlay()
        rows = [[p.get(c, "") for c in PLACE_COLUMNS] for p in po.effective_places(ov)]
        return PLACE_COLUMNS, rows

    raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown dataset")


@router.get("/export/{dataset}")
async def export_excel(dataset: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    datasets = ["users", "places", "feedback", "trips"] if dataset == "all" else [dataset]
    wb = Workbook()
    wb.remove(wb.active)

    for name in datasets:
        header, rows = await _sheet_rows(name)
        ws = wb.create_sheet(name.capitalize())
        ws.append(header)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2A44")
        ws.freeze_panes = "A2"
        for i, col in enumerate(header, start=1):
            widest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, widest + 2), 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="voyexa-{dataset}-{stamp}.xlsx"'},
    )


# --------------------------------------------------------------------------
# Settings (Power BI link)
# --------------------------------------------------------------------------
class SettingsUpdate(BaseModel):
    powerbi_url: str = ""


@router.get("/settings")
async def get_settings():
    doc = await settings_collection().find_one({"key": "powerbi_url"})
    return {"powerbi_url": (doc or {}).get("value", "")}


@router.put("/settings")
async def put_settings(body: SettingsUpdate):
    url = body.powerbi_url.strip()
    if url and not url.startswith("https://app.powerbi.com/"):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "That doesn't look like a Power BI link. It should start with https://app.powerbi.com/",
        )
    await settings_collection().update_one({"key": "powerbi_url"}, {"$set": {"key": "powerbi_url", "value": url}}, upsert=True)
    return {"powerbi_url": url}
